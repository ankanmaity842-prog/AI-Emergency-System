
import hashlib
from pathlib import Path

from pptx import Presentation
from pypdf import PdfReader

from services.embedding_service import EmbeddingService
from core.vector_store import VectorStore


class DocumentIngestionService:

    def __init__(self):
        self.embeddings = EmbeddingService()
        self.store = VectorStore()

    def load_document(self, path: Path) -> str:
        suffix = path.suffix.lower()

        if suffix in {".txt", ".md"}:
            return path.read_text(
                encoding="utf-8",
                errors="replace"
            )

        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        if suffix == ".pptx":
            presentation = Presentation(str(path))
            sections = []

            for slide_number, slide in enumerate(
                presentation.slides,
                start=1
            ):
                slide_text = []

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        text = shape.text.strip()
                        if text:
                            slide_text.append(text)

                    if shape.has_table:
                        for row in shape.table.rows:
                            cells = [
                                cell.text.strip()
                                for cell in row.cells
                            ]
                            slide_text.append(
                                " | ".join(
                                    cell for cell in cells if cell
                                )
                            )

                if slide_text:
                    sections.append(
                        f"Slide {slide_number}\n"
                        + "\n".join(slide_text)
                    )

            return "\n\n".join(sections)

        return ""

    def split_text(
        self,
        text: str,
        chunk_size: int = 1000,
        overlap: int = 150
    ) -> list[str]:
        text = " ".join(text.split())

        if not text:
            return []

        if chunk_size <= 0 or not 0 <= overlap < chunk_size:
            raise ValueError("Invalid chunk size or overlap.")

        chunks = []
        start = 0

        while start < len(text):
            end = min(start + chunk_size, len(text))
            chunks.append(text[start:end])

            if end == len(text):
                break

            start = end - overlap

        return chunks

    async def ingest_directory(
        self,
        directory: str = "data/knowledge_base"
    ) -> int:
        root = Path(directory)
        root.mkdir(parents=True, exist_ok=True)

        indexed = 0

        supported_extensions = {
            ".txt", ".md", ".pdf", ".pptx"
        }

        for path in sorted(root.rglob("*")):
            if (
                not path.is_file()
                or path.suffix.lower() not in supported_extensions
            ):
                continue

            text = self.load_document(path)
            chunks = self.split_text(text)

            if not chunks:
                continue

            source = path.name
            ids = []
            metadatas = []

            for index, chunk in enumerate(chunks):
                identifier = hashlib.sha256(
                    f"{path.resolve()}:{index}".encode()
                ).hexdigest()

                ids.append(identifier)
                metadatas.append({
                    "source": source,
                    "chunk": index
                })

            vectors = await self.embeddings.embed_texts(chunks)

            self.store.add_documents(
                ids=ids,
                documents=chunks,
                embeddings=vectors,
                metadatas=metadatas
            )

            indexed += len(chunks)

        return indexed

